from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name];s.font.name='Microsoft YaHei';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.28
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(13)
head=sec.header.paragraphs[0];head.text='ECHO CAMPUS     使用与场景替换指南';head.style='Caption'
foot=sec.footer.paragraphs[0];foot.text='MeetMind  |  原型交付  |  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,2)
def page(title):doc.add_page_break();doc.add_heading(title,1)
def steps(rows):
 for i,t in enumerate(rows,1):p(f'{i}.  {t}')
doc.add_heading('Echo Campus 使用与场景替换指南',0)
p('用于合作展示、双端体验与下一次场景试验','Subtitle')
p('这份指南帮助你打开原型、向合作方演示体验，并替换为手工 GLB 或 Marble 导出的场景。当前版本已完成桌面与移动视口的浏览器画面和主要交互验证；实体手机、NFC 标签及真实 Marble 导出文件需在活动前联调。')
doc.add_picture(str(ROOT/'assets/campus.png'),width=Inches(6.9))
p('白庭校园的同源程序化几何离线参考图。它用于评审空间与构图，不替代浏览器实机效果。','Caption')
h('三个入口')
p('互动原型  https://capture.meetmind.online/echo-campus/')
p('大屏模式  https://capture.meetmind.online/echo-campus/?mode=stage')
p('独立演示窗口  https://capture.meetmind.online/echo-campus/?demoSession=tab')
p('公开演示使用虚构昵称即可。请先检查在线场景，再向合作方发送展示视频和合作方案。')

page('三分钟展示流程')
steps(['打开互动原型，先停留在白庭校园全景。拖动画面、滚轮缩放，切换全景、入口、庭院和俯瞰，展示空间品质。','用第一台设备领取分身，确认昵称、身份、我能提供、我想认识与公开资料同意项。分身会进入同一个活动世界。','同一台电脑可在新窗口打开 ?demoSession=tab，领取另一位分身；第二台设备也可用普通入口。普通标签页会共享身份，独立演示窗口使用单独会话。','查看自己的名片并打开推荐，展开“为什么推荐这次相遇”，展示供给和需求原文。','点选另一位真实领取的参与者，发出相遇请求。到接收方的“我的相遇”中确认，观察公共连接更新。','切换水上艺廊，说明人物名片和已确认关系保留，建筑与景观可以按活动主题重做。'])
h('演示时应说清楚')
p('内置人物和关系是虚构示例；内置人物没有会话，不会替它们自动确认邀请。双向确认必须使用两个真实领取的演示会话。')
p('分身是程序化风格化人物，不是照片重建或真人克隆。推荐来自公开供需标签匹配，不是大模型推理；解释可以对照原文。')
h('NFC 的现场接法')
p('把同一个 HTTPS 入口写入 NFC 标签，手机触碰后打开网页。芯片写入、手机兼容、浏览器权限与现场网络需用最终设备联调。公开标签负责打开入口；真实活动的身份核验和凭据应另行设计。')

page('明天替换 GLB 或 Marble 场景')
steps(['先准备一个允许用于本次展示的自包含 GLB，或 Marble 导出的 SPZ。优先从较小文件开始。','右上角打开“场景”，选择本地文件。文件在本浏览器读取，不会上传到活动服务。','先调方向，再调尺度，最后调位置。GLB 通常 Y 轴向上；SPZ 初值 X=180° 只是常见轴系参考，必须根据实际画面判断。','以人物身高和门高作尺度参照。把可活动地面放到接近 y=0 的位置，再设置地面高度、出生点和人物站位。','在高级 JSON 中设置 bounds、anchors、cameras 与 colliders。模型变换只影响视觉，活动坐标需要独立对齐。','点击应用，切换四个镜头检查构图与人物尺度。导入失败应保留旧场景；不要把错误提示忽略为已经换景。','保存本机配置或导出 JSON。将 JSON 和模型文件一起保留；本地模型本身不会被浏览器设置永久保存。'])
h('必须知道的限制')
p('本地文件上限 512 MB；这不是流畅运行的性能承诺。较大的 Splat 和贴图应在最终桌面、手机上分别验证。')
p('本地 GLTF 必须内嵌资源；多个外链纹理和二进制文件建议先合并导出为 GLB。当前未配置 Draco 和 Meshopt 专用解码器。')
p('通用导入采用活动平面和圆形碰撞壳，不会从 Splat 自动恢复楼梯、坡道或精确墙体。需要这些能力的场景应补充地形适配。')

page('让不同设备进入同一个新场景')
p('本地导入只改变当前浏览器。若要让手机与大屏看到同一建筑，需要把模型与配置放到所有设备可读取的地址。')
steps(['把模型放入 public/scenes/，将配置中的 url 写为相对配置文件的路径，例如 ./my-campus.glb。','配置使用 schema echo-campus.scene.v1。确定 type、模型变换、groundY、bounds、spawn、anchors 和 cameras 后保存为 my-campus.json。','分享入口 https://capture.meetmind.online/echo-campus/?sceneManifest=./scenes/my-campus.json 。模型地址会相对这份 JSON 解析。','若要全活动默认进入此场景，编辑 public/scene-startup.json，设 enabled 为 true，并在 manifest 中填写完整配置。构建后发布。','重新在两台设备打开入口，分别确认建筑、人物站位、入口与已确认关系。'])
h('自检样本')
p('GLB 白展厅  ?sceneManifest=./scenes/import-test.glb.json')
p('自造 SPZ  ?sceneManifest=./scenes/import-test.spz.json')
p('这些样本用于检查导入路径。GLB 与自造 SPZ 均已在真实浏览器通过共享启动配置加载并呈现。自造 SPZ 不是 Marble 生成案例；每份真实 Marble 导出文件仍需校准方向、尺度与活动层。')
h('手写第三套场景')
p('在 src/scenes/ 新建生成函数，返回 root、bounds、spawn、anchors、cameras、colliders、update 和 dispose。然后在 src/runtime/SceneRegistry.js 登记唯一 id、displayName、helper 与 factory。场景卡自动出现在选择器。')
p('完整字段示例与生成函数契约见随包的 SCENE-SWAP.md。')

page('运行维护与交付边界')
h('本地运行')
p('进入 showcase/echo-campus 后依次执行 npm ci、npm run build、npm start；打开 http://127.0.0.1:5189/。需要改代码热更新时另开 npm run dev，默认端口 5190。')
h('当前独立部署')
p('代码位于 /root/meetmind_wt_main/showcase/echo-campus。systemd 服务 echo-campus-showcase 监听 127.0.0.1:5191，Nginx 将 /echo-campus/ 转发到该服务。演示数据保存在 /var/lib/echo-campus/event.json。')
p('现有 EchoWorld 和其数据目录不受影响。不要把此应用构建覆盖到 /var/www/echoworld，也不要在共享 /root/meetmind_go 执行本项目构建或 Git 操作。')
h('已验证与尚需验证')
p('截至本版，生产构建通过，50 项测试全部通过，覆盖 HTTP/WS 双会话、资料同意、身份恢复、请求确认、浏览器 fetch、独立演示会话、错误恢复、导入配置及共享场景启动。')
p('真实浏览器已验证 1440×900 桌面与 492×898 移动视口的画面、镜头、名片、推荐、GLB 与自造 SPZ 加载。两个独立会话完成“周澈发起、林予确认”，双方均显示 1 条已点亮连接；校园切换为艺廊后，身份与连接保留。实体手机、NFC 硬件和真实 Marble 导出文件尚未验收。')
h('随包内容')
p('展示主片 Echo-Campus-Showcase.mp4：97 秒、1080p、30 帧，采用实际 3D 场景画布录屏与真实浏览器界面截图剪辑，配 DashScope 中文旁白和字幕；不是全程交互录屏。')
p('主片链接  https://capture.meetmind.online/echo-campus/Echo-Campus-Showcase.mp4')
p('旧片 Echo-Campus-97s-preview.mp4 保留为建筑参考与流程示意预览。另附合作方案 PDF、可运行源码、换场景样本、本指南及 Markdown 配置说明。')

# Preserve the verified Word layout and accessible image description.
for style_name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Header','Footer']:
 try:
  style=doc.styles[style_name];style.font.color.rgb=RGBColor(0,0,0)
  prop=style._element.find(qn('w:pPr'))
  if prop is not None:
   border=prop.find(qn('w:pBdr'))
   if border is not None:prop.remove(border)
 except KeyError:pass
for para in doc.paragraphs:
 prop=para._p.get_or_add_pPr();border=prop.find(qn('w:pBdr'))
 if border is not None:prop.remove(border)
 for run in para.runs:run.font.color.rgb=RGBColor(0,0,0)
for section in doc.sections:
 for container in [section.header,section.footer]:
  for para in container.paragraphs:
   for run in para.runs:run.font.color.rgb=RGBColor(0,0,0)
for shape in doc.inline_shapes:
 shape._inline.docPr.set('descr','白庭校园的白色弧形建筑、水庭、连廊与景观示意图')
 shape._inline.docPr.set('title','白庭校园空间示意')

doc.save(ROOT/'Echo-Campus-使用与换场景指南.docx')
print('created docx')
